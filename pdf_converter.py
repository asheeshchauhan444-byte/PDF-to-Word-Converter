import fitz
from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.enum.section import WD_SECTION

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
OUTPUT_DIR = BASE_DIR / "outputs"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

def convert_pdf_to_word(pdf_path, output_path, dpi=200):
    pdf_path = Path(pdf_path)
    output_path = Path(output_path)
    if not pdf_path.exists():
        raise FileNotFoundError(f"PDF not found: {pdf_path}")
    pdf = fitz.open(str(pdf_path))
    document = Document()
    first_page = True
    for page_index, page in enumerate(pdf):
        zoom = dpi / 72
        matrix = fitz.Matrix(zoom, zoom)
        pixmap = page.get_pixmap(matrix=matrix, alpha=False)
        image_path = OUTPUT_DIR / f"_page_{page_index + 1}.png"
        pixmap.save(str(image_path))
        page_width = page.rect.width
        page_height = page.rect.height
        if first_page:
            section = document.sections[0]
            first_page = False
        else:
            section = document.add_section(WD_SECTION.NEW_PAGE)
        width_inches = page_width / 72
        height_inches = page_height / 72
        section.page_width = Inches(width_inches)
        section.page_height = Inches(height_inches)
        section.top_margin = Inches(0)
        section.bottom_margin = Inches(0)
        section.left_margin = Inches(0)
        section.right_margin = Inches(0)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = 0
        paragraph.paragraph_format.space_after = 0
        run = paragraph.add_run()
        run.add_picture(str(image_path), width=Inches(width_inches), height=Inches(height_inches))
    document.save(str(output_path))
    for temporary_file in OUTPUT_DIR.glob("_page_*.png"):
        try:
            temporary_file.unlink()
        except Exception:
            pass
    pdf.close()
    return output_path
